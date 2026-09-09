import csv
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from response_parser import filter_responses


class ResponseTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.output = self.directory / "filtered.xlsx"

    def write(self, name, rows):
        path = self.directory / name
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            csv.writer(handle).writerows(rows)
        return path

    def result(self):
        from openpyxl import load_workbook
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames, ["Responses"])
            self.assertTrue(all(cell.data_type != "f" for row in book.active for cell in row))
            return [[value if value is not None else "" for value in row]
                    for row in book.active.iter_rows(values_only=True)]
        finally:
            book.close()

    def test_membership_preserves_response_data_and_duplicates(self):
        header = ["Answer", " Computing ID ", "Answer", ""]
        kept = [[" None ", " JS1 ", "line one\nline two", "001"],
                ["=1+1", "js1", "a,b", ""]]
        responses = self.write("responses.csv", [header, kept[0],
            ["excluded", "other"], ["blank", ""], ["null", "None"], kept[1]])
        consented = self.write("consented.csv", [["computing id", "IRBs"],
            ["js1", "A"], ["JS1", "B"], ["", "A"], ["None", "A"]])
        self.assertEqual(filter_responses(responses, consented, self.output), 2)
        self.assertEqual(self.result(), [header, *kept])
        with self.assertRaises(FileExistsError):
            filter_responses(responses, consented, self.output)
        self.assertEqual(self.result(), [header, *kept])

    def test_empty_allowlist_keeps_header(self):
        responses = self.write("responses.csv", [["Computing ID", "Q"], ["a", "Yes"]])
        consented = self.write("consented.csv", [["Computing ID"]])
        self.assertEqual(filter_responses(responses, consented, self.output), 0)
        self.assertEqual(self.result(), [["Computing ID", "Q"]])

    def test_invalid_headers_leave_no_output(self):
        consented = self.write("consented.csv", [["Computing ID"], ["a"]])
        for rows in ([], [["Email"]], [["Computing ID", " COMPUTING ID "]]):
            responses = self.write("responses.csv", rows)
            with self.assertRaises(ValueError):
                filter_responses(responses, consented, self.output)
            self.assertFalse(self.output.exists())

    def test_excel_sheets_headers_and_formula_rejection(self):
        from openpyxl import Workbook
        path = self.directory / "responses.xlsx"
        book = Workbook()
        sheet = book.create_sheet("Responses")
        sheet.append(["Survey title"])
        sheet.append(["Computing ID", "Q"])
        sheet.append([12, " None "])
        sheet["A3"].number_format = "0000"
        ids = book.create_sheet("Allowed")
        ids.append(["Computing ID"])
        ids.append(["0012"])
        book.save(path)
        options = dict(response_sheet="Responses", consent_sheet="Allowed", response_header_row=2)
        self.assertEqual(filter_responses(path, path, self.output, **options), 1)
        self.assertEqual(self.result(), [["Computing ID", "Q"], ["0012", " None "]])
        sheet["B3"] = "=1+1"
        book.save(path)
        book.close()
        failed_output = self.directory / "failed.xlsx"
        with self.assertRaisesRegex(ValueError, "Formula"):
            filter_responses(path, path, failed_output, **options)
        self.assertFalse(failed_output.exists())


if __name__ == "__main__":
    unittest.main()

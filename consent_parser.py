"""Filter any number of IRB sheets and merge consenting people by Computing ID."""
from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
import json
from pathlib import Path
import re
from typing import Iterable

COLUMNS = ("First name", "Last name", "email", "Computing ID", "Subject ID", "Professor", "Consent")
OUTPUT_COLUMNS = (*COLUMNS[:-1], "IRBs")


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalized(value: object) -> str:
    return " ".join(clean(value).split()).casefold()


@dataclass
class IRB:
    name: str
    file: str | Path
    consent_values: list[str]
    sheet: str | None = None
    header_row: int = 1


def _cell_text(cell: object) -> str:
    # Preserve common zero-padded numeric identifiers, e.g. format 00000.
    value = cell.value
    if cell.data_type == "f":
        raise ValueError("Formula cells are unsupported; export a values-only copy first.")
    if cell.data_type == "e":
        raise ValueError("An Excel error cell was encountered.")
    fmt = cell.number_format
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value).is_integer():
            integer = int(value)
            if re.fullmatch(r"0+", fmt or "") and integer >= 0:
                return str(integer).zfill(len(fmt))
            return str(integer)
    return clean(value)


def read_rows(path: str | Path, sheet: str | None = None) -> Iterable[list[str]]:
    """Stream .xlsx/.xlsm or UTF-8 .csv rows; default to the first worksheet."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        if sheet is not None:
            raise ValueError("Worksheet names cannot be used with CSV files.")
        with path.open(newline="", encoding="utf-8-sig") as handle:
            for row in csv.reader(handle):
                yield [clean(value) for value in row]
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is not None and sheet not in book.sheetnames:
                raise ValueError(f"Worksheet {sheet!r} does not exist in {path.name}.")
            worksheet = book[sheet] if sheet is not None else book.worksheets[0]
            for row in worksheet.iter_rows():
                yield [_cell_text(cell) for cell in row]
        finally:
            book.close()
    else:
        raise ValueError(f"Unsupported input extension {suffix!r}; use .xlsx, .xlsm, or .csv.")


def parse_irb(irb: IRB) -> list[dict[str, str]]:
    """Read the common schema and keep only explicitly accepted consent values."""
    if not isinstance(irb.consent_values, list) or not irb.consent_values or any(
        not isinstance(value, str) or not normalized(value) for value in irb.consent_values
    ):
        raise ValueError(f"{irb.name}: consent_values must be a nonempty list of nonblank strings.")
    if not isinstance(irb.header_row, int) or irb.header_row < 1:
        raise ValueError(f"{irb.name}: header_row must be a positive integer.")
    accepted = {normalized(value) for value in irb.consent_values}
    rows = read_rows(irb.file, irb.sheet)
    try:
        headers = None
        for _ in range(irb.header_row):
            headers = next(rows, None)
        if headers is None:
            raise ValueError(f"{irb.name}: missing header row.")
        headers = [normalized(header) for header in headers]
        nonblank = [header for header in headers if header]
        if len(nonblank) != len(set(nonblank)):
            raise ValueError(f"{irb.name}: duplicate headers.")
        missing = [column for column in COLUMNS if normalized(column) not in headers]
        if missing:
            raise ValueError(f"{irb.name}: missing columns: {', '.join(missing)}")
        indices = {column: headers.index(normalized(column)) for column in COLUMNS}
        result = []
        for number, row in enumerate(rows, irb.header_row + 1):
            record = {column: clean(row[index]) if index < len(row) else ""
                      for column, index in indices.items()}
            if normalized(record["Consent"]) not in accepted:
                continue
            if not record["Computing ID"]:
                raise ValueError(f"{irb.name}, row {number}: consenting row has no Computing ID.")
            result.append(record)
        return result
    finally:
        rows.close()


def parse_consent_files(irbs: Iterable[IRB]) -> list[dict[str, str]]:
    """Include anyone consenting to at least one IRB; merge only by Computing ID."""
    people = {}
    memberships = {}
    names = set()
    for irb in irbs:
        name = clean(irb.name)
        if not name or normalized(name) in names:
            raise ValueError("Each IRB needs a unique, nonblank name.")
        names.add(normalized(name))
        for record in parse_irb(irb):
            key = normalized(record["Computing ID"])
            if key not in people:
                people[key] = {column: "" for column in OUTPUT_COLUMNS}
                memberships[key] = []
            person = people[key]
            for column in COLUMNS[:-1]:
                value = record[column]
                if value and person[column] and normalized(value) != normalized(person[column]):
                    raise ValueError(
                        f"{name}: conflicting {column} for Computing ID {key!r} "
                        f"(previous IRBs: {', '.join(memberships[key])})."
                    )
                if value and not person[column]:
                    person[column] = value
            if name not in memberships[key]:
                memberships[key].append(name)
            person["IRBs"] = "; ".join(memberships[key])
    return sorted(people.values(), key=lambda row: (
        normalized(row["Last name"]), normalized(row["First name"]), normalized(row["Computing ID"])
    ))


def load_irbs(path: str | Path) -> list[IRB]:
    """Load a JSON list; resolve input files relative to the configuration file."""
    path = Path(path)
    with path.open(encoding="utf-8-sig") as handle:
        config = json.load(handle)
    if not isinstance(config, list) or not config:
        raise ValueError("Configuration must be a nonempty list of IRBs.")
    irbs = []
    for entry in config:
        if not isinstance(entry, dict):
            raise ValueError("Each IRB configuration must be an object.")
        irb = IRB(**entry)
        irb.file = path.parent / irb.file
        irbs.append(irb)
    return irbs


def write_csv(rows: Iterable[dict[str, str]], path: str | Path) -> None:
    """Write an Excel-compatible CSV without overwriting an existing file."""
    with Path(path).open("x", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("config", type=Path, help="JSON file defining the IRBs")
    parser.add_argument("-o", "--output", type=Path, default=Path("consented.csv"))
    args = parser.parse_args()
    try:
        if args.output.suffix.lower() != ".csv":
            raise ValueError("Output must have a .csv extension.")
        rows = parse_consent_files(load_irbs(args.config))
        write_csv(rows, args.output)
    except (OSError, ValueError, TypeError, ImportError) as exc:
        parser.exit(2, f"Error: {exc}\n")
    print(f"Wrote {len(rows)} participants to {args.output}")


if __name__ == "__main__":
    main()

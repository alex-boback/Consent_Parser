"""Filter response data by Computing IDs in a consented participant file."""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from consent_parser import _cell_text, normalized


def read_rows(path: str | Path, sheet: str | None = None) -> list[list[str]]:
    """Read values without trimming response text or replacing literal 'None'."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        if sheet is not None:
            raise ValueError("Worksheet names cannot be used with CSV files.")
        with path.open(newline="", encoding="utf-8-sig") as handle:
            return list(csv.reader(handle))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            worksheet = book[sheet] if sheet is not None else book.worksheets[0]
            rows = []
            for row in worksheet.iter_rows():
                values = []
                for cell in row:
                    # Validate formulas/errors and retain numeric ID formatting.
                    value = _cell_text(cell)
                    values.append(cell.value if isinstance(cell.value, str) else value)
                rows.append(values)
            return rows
        finally:
            book.close()
    else:
        raise ValueError("Unsupported input extension; use .xlsx, .xlsm, or .csv.")


def read_table(path, sheet=None, header_row=1):
    """Return the original header and each data row paired with its normalized ID."""
    rows = read_rows(path, sheet)
    if not 1 <= header_row <= len(rows):
        raise ValueError(f"{path}: header row is outside the file.")
    header = rows[header_row - 1]
    columns = [normalized(value) for value in header]
    if columns.count("computing id") != 1:
        raise ValueError(f"{path}: expected exactly one Computing ID column.")
    index = columns.index("computing id")

    # Short CSV rows may omit the ID column; treat those IDs as blank.
    data = [(normalized(row[index]) if index < len(row) else "", row)
            for row in rows[header_row:]]
    return header, data


def filter_responses(
    responses: str | Path,
    consented: str | Path,
    output: str | Path,
    *,
    response_sheet: str | None = None,
    consent_sheet: str | None = None,
    response_header_row: int = 1,
    consent_header_row: int = 1,
) -> int:
    """Write matching rows to a new Excel file and return their count."""
    from openpyxl import Workbook

    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have an .xlsx extension.")

    # A set makes membership lookups fast and excludes missing consent IDs.
    _, consent_rows = read_table(consented, consent_sheet, consent_header_row)
    allowed = {identifier for identifier, _ in consent_rows if identifier}

    # Filter whole rows, preserving response text, order, and repeated responses.
    header, response_rows = read_table(responses, response_sheet, response_header_row)
    selected = [row for identifier, row in response_rows if identifier in allowed]

    book = Workbook()
    sheet = book.active
    sheet.title = "Responses"
    for row_number, row in enumerate([header, *selected], start=1):
        for column, value in enumerate(row, start=1):
            # Keep IDs and response text literal, including leading zeros and '='.
            cell = sheet.cell(row_number, column, value)
            cell.data_type = "s"

    # Binary exclusive mode prevents overwriting an existing file.
    try:
        with output.open("xb") as handle:
            book.save(handle)
    finally:
        book.close()
    return len(selected)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("responses", type=Path, help="Response CSV or Excel file")
    parser.add_argument("consented", type=Path, help="Existing parser output or other ID list")
    parser.add_argument("-o", "--output", type=Path, default=Path("filtered_responses.xlsx"))
    parser.add_argument("--response-sheet")
    parser.add_argument("--consent-sheet")
    parser.add_argument("--response-header-row", type=int, default=1)
    parser.add_argument("--consent-header-row", type=int, default=1)
    args = parser.parse_args()
    try:
        count = filter_responses(**vars(args))
    except (OSError, ValueError, KeyError, ImportError) as exc:
        parser.exit(2, f"Error: {exc}\n")
    print(f"Wrote {count} response rows to {args.output}")


if __name__ == "__main__":
    main()
